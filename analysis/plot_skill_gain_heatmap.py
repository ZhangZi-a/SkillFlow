#!/usr/bin/env python3
"""Plot a task-family-by-agent heatmap for skill completion-rate gains.

The script scans top-level job summaries under ``jobs/*/summary.json``, matches
base/skill pairs, maps task IDs to task families from the Excel lookup table,
and renders a heatmap whose cells are the average
``completion_rate(skill) - completion_rate(base)`` within each family.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_JOBS_DIR = ROOT / "jobs"
DEFAULT_OUTPUT_DIR = ROOT / "analysis" / "figures"
DEFAULT_WORKFLOW_XLSX = ROOT / "数据分析与分类.xlsx"
DEFAULT_BASENAME = "jobs_skill_gain_heatmap_completion_delta"
DEFAULT_VMAX = 0.4
FONT_SCALE = 1.8
AXIS_TITLE_FONT_SIZE = 30.0
CELL_SIZE_IN = 0.8
FIG_MARGIN_X = 4.2
FIG_MARGIN_Y = 3.4
CELL_BORDER_WIDTH = 1.1

AGENT_MODEL_LABEL_OVERRIDES = {
    "claude-code-minimax2dot5": "Claude Code + MiniMax 2.5",
    "claude-code-minimax2dot7": "Claude Code + MiniMax 2.7",
    "claude-code-opus4dot5": "Claude Code + Opus 4.5",
    "claude-code-opus4dot6": "Claude Code + Opus 4.6",
    "claude-code-sonnet4dot5": "Claude Code + Sonnet 4.5",
    "claude-code-sonnet4dot6": "Claude Code + Sonnet 4.6",
    "codex-cli-gpt-5.3-codex": "Codex CLI + GPT-5.3-Codex",
    "codex-cli-gpt-5.4": "Codex CLI + GPT-5.4",
    "kimi-cli-kimi-k2dot5": "Kimi CLI + Kimi-K2.5",
    "qwen-coder-qwen3-coder-480b-a35b-instruct": "Qwen Coder + Qwen3-480B-A35B",
    "qwen-coder-qwen3-coder-mext": "Qwen Coder + Qwen3-Mext",
}

FAMILY_LABEL_FALLBACKS = {
    "econ-detrending-correlation": "Finance & Economics",
    "harbor_gdpval_20": "Finance & Economics",
    "sec-financial-report": "Finance & Economics",
    "harbor_gdpval_21": "Operations & Supply Chain",
    "harbor_gdpval_36": "Operations & Supply Chain",
    "merge_20_21": "Operations & Supply Chain",
    "merge_35_37": "Operations & Supply Chain",
    "merge_36_41": "Operations & Supply Chain",
    "harbor_gdpval_42": "Healthcare & Life Sciences",
    "lab-unit-harmonization": "Healthcare & Life Sciences",
    "harbor_gdpval_3": "Governance & Strategy",
    "harbor_gdpval_33": "Governance & Strategy",
    "invoice-fraud-detection": "Governance & Strategy",
    "exceltable-in-ppt": "Data & Document Intelligence",
    "jpg-ocr-stat": "Data & Document Intelligence",
    "merge_court_offer": "Data & Document Intelligence",
    "merge_pdf_xlsx": "Data & Document Intelligence",
    "merge_weight_reserves": "Data & Document Intelligence",
    "pptx-reference-formatting": "Data & Document Intelligence",
    "sales-pivot-analysis": "Data & Document Intelligence",
}

FAMILY_ORDER = [
    "Finance & Economics",
    "Operations & Supply Chain",
    "Healthcare & Life Sciences",
    "Governance & Strategy",
    "Data & Document Intelligence",
]

FAMILY_LABEL_ABBREVIATIONS = {
    "Finance & Economics": "FE",
    "Operations & Supply Chain": "OS",
    "Healthcare & Life Sciences": "HLS",
    "Governance & Strategy": "GS",
    "Data & Document Intelligence": "DDI",
}


@dataclass(frozen=True)
class PairSummary:
    base_slug: str
    skill_slug: str
    agent_model_label: str
    task_order: list[str]
    base_rates: dict[str, float]
    skill_rates: dict[str, float]
    avg_delta: float


@dataclass(frozen=True)
class CellValue:
    family_label: str
    agent_model_label: str
    base_slug: str
    skill_slug: str
    n_tasks: int
    delta: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Plot a skill-gain heatmap from Harbor job summaries.")
    parser.add_argument(
        "--jobs-dir",
        type=Path,
        default=DEFAULT_JOBS_DIR,
        help=f"Directory containing top-level jobs folders (default: {DEFAULT_JOBS_DIR})",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for figure outputs (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument(
        "--workflow-xlsx",
        type=Path,
        default=DEFAULT_WORKFLOW_XLSX,
        help=f"Excel lookup table for workflow families (default: {DEFAULT_WORKFLOW_XLSX})",
    )
    parser.add_argument(
        "--basename",
        default=DEFAULT_BASENAME,
        help="Base filename for generated heatmap files.",
    )
    parser.add_argument(
        "--vmax",
        type=float,
        default=DEFAULT_VMAX,
        help="Symmetric color limit for completion deltas (default: 0.4).",
    )
    return parser.parse_args()


def load_json(path: Path) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        msg = f"Expected a JSON object in {path}"
        raise ValueError(msg)
    return data


def is_skill_slug(slug: str) -> bool:
    return slug.endswith("-skill") or slug.endswith("-w-skill")


def infer_base_slug(skill_slug: str) -> str | None:
    if skill_slug.endswith("-w-skill"):
        return skill_slug.removesuffix("-w-skill") + "-wo-skill"
    if skill_slug.endswith("-skill"):
        return skill_slug.removesuffix("-skill")
    return None


def _generic_agent_model_label(slug: str) -> str:
    label = slug.removesuffix("-skill")
    label = label.replace("claude-code-", "Claude Code + ")
    label = label.replace("codex-cli-", "Codex CLI + ")
    label = label.replace("kimi-cli-", "Kimi CLI + ")
    label = label.replace("qwen-coder-", "Qwen Coder + ")
    label = label.replace("dot", ".")
    label = label.replace("gpt-5.3-codex", "GPT-5.3-Codex")
    label = label.replace("gpt-5.4", "GPT-5.4")
    label = label.replace("kimi-k2.5", "Kimi-K2.5")
    label = label.replace("qwen3-coder-480b-a35b-instruct", "Qwen3-480B-A35B")
    label = label.replace("qwen3-coder-mext", "Qwen3-Mext")
    label = label.replace("opus4.5", "Opus 4.5")
    label = label.replace("opus4.6", "Opus 4.6")
    label = label.replace("sonnet4.5", "Sonnet 4.5")
    label = label.replace("sonnet4.6", "Sonnet 4.6")
    label = label.replace("minimax2.5", "MiniMax 2.5")
    label = label.replace("minimax2.7", "MiniMax 2.7")
    return " ".join(label.split()).strip("+- ")


def agent_model_label_from_slug(slug: str) -> str:
    if slug in AGENT_MODEL_LABEL_OVERRIDES:
        return AGENT_MODEL_LABEL_OVERRIDES[slug]
    return _generic_agent_model_label(slug)


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def pick_english_label(value: Any) -> str | None:
    if value is None:
        return None
    parts = [part.strip() for part in str(value).splitlines() if str(part).strip()]
    if not parts:
        return None
    for part in reversed(parts):
        if any(("A" <= ch <= "Z") or ("a" <= ch <= "z") for ch in part):
            return part
    return parts[-1]


def load_task_family_map(xlsx_path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    if not xlsx_path.is_file():
        return mapping
    try:
        import openpyxl
    except Exception:
        return mapping

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return mapping

    headers = [normalize_header(value) for value in rows[0]]
    workflow_name_index = next((i for i, h in enumerate(headers) if "workflow name" in h), None)
    core_category_index = next((i for i, h in enumerate(headers) if "core category" in h), None)
    if workflow_name_index is None or core_category_index is None:
        return mapping

    current_family: str | None = None
    for row in rows[1:]:
        raw_task_name = row[workflow_name_index]
        raw_family = row[core_category_index]
        if raw_family is not None:
            current_family = pick_english_label(raw_family)
        if raw_task_name is None or not current_family:
            continue
        task_name = str(raw_task_name).strip()
        if task_name:
            mapping[task_name] = current_family
    return mapping


def task_family_label(task_name: str, family_map: dict[str, str]) -> str:
    return family_map.get(task_name, FAMILY_LABEL_FALLBACKS.get(task_name, "Other"))


def wrap_agent_model_label(label: str) -> str:
    if " + " in label:
        left, right = label.split(" + ", maxsplit=1)
        return f"{left}\n{right}"
    return label


def family_rank(slug: str) -> tuple[int, str]:
    lowered = slug.lower()
    if "claude" in lowered:
        return (0, slug)
    if "codex" in lowered:
        return (1, slug)
    if "kimi" in lowered:
        return (2, slug)
    if "qwen" in lowered:
        return (3, slug)
    return (9, slug)


def load_top_level_summaries(jobs_dir: Path) -> dict[str, dict[str, Any]]:
    summaries: dict[str, dict[str, Any]] = {}
    for child in sorted(jobs_dir.iterdir()):
        if not child.is_dir():
            continue
        summary_path = child / "summary.json"
        if not summary_path.is_file():
            continue
        summaries[child.name] = load_json(summary_path)
    if not summaries:
        msg = f"No top-level summary.json found under {jobs_dir}"
        raise FileNotFoundError(msg)
    return summaries


def group_completion_rates(summary: dict[str, Any]) -> tuple[list[str], dict[str, float]]:
    groups = summary.get("groups")
    if not isinstance(groups, list):
        msg = "Summary JSON missing groups list"
        raise ValueError(msg)
    order: list[str] = []
    rates: dict[str, float] = {}
    for group in groups:
        if not isinstance(group, dict):
            continue
        name = group.get("name")
        completion_rate = group.get("completion_rate")
        if not isinstance(name, str) or completion_rate is None:
            continue
        order.append(name)
        rates[name] = float(completion_rate)
    return order, rates


def average_delta(base_rates: dict[str, float], skill_rates: dict[str, float], task_order: list[str]) -> float:
    deltas = [skill_rates[task] - base_rates[task] for task in task_order if task in base_rates and task in skill_rates]
    return sum(deltas) / len(deltas) if deltas else float("nan")


def build_pairs(jobs_dir: Path) -> list[PairSummary]:
    summaries = load_top_level_summaries(jobs_dir)
    pairs: list[PairSummary] = []
    for skill_slug in sorted(summaries, key=family_rank):
        if not is_skill_slug(skill_slug):
            continue
        base_slug = infer_base_slug(skill_slug)
        if base_slug is None or base_slug not in summaries:
            continue
        base_order, base_rates = group_completion_rates(summaries[base_slug])
        skill_order, skill_rates = group_completion_rates(summaries[skill_slug])
        task_order = [task for task in base_order if task in skill_rates]
        if not task_order:
            continue
        if set(task_order) != set(skill_order):
            missing = sorted(set(skill_order) - set(task_order))
            for task in missing:
                task_order.append(task)
        pairs.append(
            PairSummary(
                base_slug=base_slug,
                skill_slug=skill_slug,
                agent_model_label=agent_model_label_from_slug(base_slug),
                task_order=task_order,
                base_rates=base_rates,
                skill_rates=skill_rates,
                avg_delta=average_delta(base_rates, skill_rates, task_order),
            )
        )
    if not pairs:
        msg = f"No matched base/skill summary pairs found under {jobs_dir}"
        raise ValueError(msg)
    return sorted(pairs, key=lambda pair: (*family_rank(pair.base_slug), -pair.avg_delta, pair.agent_model_label))


def aggregate_family_deltas(pair: PairSummary, family_map: dict[str, str]) -> dict[str, tuple[float, int]]:
    grouped: dict[str, list[float]] = {}
    for task_name in pair.task_order:
        base_rate = pair.base_rates.get(task_name)
        skill_rate = pair.skill_rates.get(task_name)
        if base_rate is None or skill_rate is None:
            continue
        family_label = task_family_label(task_name, family_map)
        grouped.setdefault(family_label, []).append(skill_rate - base_rate)

    aggregated: dict[str, tuple[float, int]] = {}
    for family_label, deltas in grouped.items():
        aggregated[family_label] = (sum(deltas) / len(deltas), len(deltas))
    return aggregated


def ordered_family_labels(available_families: set[str]) -> list[str]:
    labels = [family for family in FAMILY_ORDER if family in available_families]
    remaining = sorted(available_families - set(labels))
    return labels + remaining


def build_matrix(
    pairs: list[PairSummary],
    family_map: dict[str, str],
) -> tuple[list[str], list[str], list[list[float]], list[CellValue]]:
    pair_family_values = [aggregate_family_deltas(pair, family_map) for pair in pairs]
    family_labels = ordered_family_labels({family for values in pair_family_values for family in values})
    agent_model_labels = [pair.agent_model_label for pair in pairs]

    matrix: list[list[float]] = []
    cells: list[CellValue] = []
    for family_label in family_labels:
        row: list[float] = []
        for pair, family_values in zip(pairs, pair_family_values):
            delta, n_tasks = family_values.get(family_label, (math.nan, 0))
            row.append(delta)
            cells.append(
                CellValue(
                    family_label=family_label,
                    agent_model_label=pair.agent_model_label,
                    base_slug=pair.base_slug,
                    skill_slug=pair.skill_slug,
                    n_tasks=n_tasks,
                    delta=delta,
                )
            )
        matrix.append(row)
    return family_labels, agent_model_labels, matrix, cells


def render_heatmap(
    family_labels: list[str],
    agent_model_labels: list[str],
    matrix: list[list[float]],
    output_paths: list[Path],
    vmax: float,
) -> list[Path]:
    try:
        import matplotlib.pyplot as plt
        import numpy as np
        from mpl_toolkits.axes_grid1 import make_axes_locatable
    except Exception as exc:
        msg = "matplotlib and numpy are required to render the heatmap"
        raise RuntimeError(msg) from exc

    data = np.array(matrix, dtype=float)
    masked = np.ma.masked_invalid(data)

    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "svg.fonttype": "none",
            "font.size": 10 * FONT_SCALE,
            "axes.labelsize": 11 * FONT_SCALE,
            "xtick.labelsize": 9,
            "ytick.labelsize": 10 * FONT_SCALE,
        }
    )

    fig_width = max(12.0, CELL_SIZE_IN * len(agent_model_labels) + FIG_MARGIN_X)
    fig_height = max(7.8, CELL_SIZE_IN * len(family_labels) + FIG_MARGIN_Y)
    fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=240)

    cmap = plt.get_cmap("RdBu_r").copy()
    cmap.set_bad("#F3F4F6")
    image = ax.imshow(
        masked,
        cmap=cmap,
        vmin=-vmax,
        vmax=vmax,
        aspect="equal",
        interpolation="nearest",
    )

    ax.set_xticks(range(len(agent_model_labels)))
    ax.set_xticklabels(
        [wrap_agent_model_label(label) for label in agent_model_labels],
        rotation=18,
        ha="right",
        rotation_mode="anchor",
    )
    ax.set_yticks(range(len(family_labels)))
    ax.set_yticklabels([FAMILY_LABEL_ABBREVIATIONS.get(label, label) for label in family_labels])
    ax.tick_params(axis="x", length=0, pad=10)
    ax.tick_params(axis="y", length=0, pad=8)

    ax.set_xticks([index - 0.5 for index in range(len(agent_model_labels) + 1)], minor=True)
    ax.set_yticks([index - 0.5 for index in range(len(family_labels) + 1)], minor=True)
    ax.grid(which="minor", color="white", linewidth=CELL_BORDER_WIDTH)
    ax.tick_params(which="minor", bottom=False, left=False)

    for row_index, row in enumerate(matrix):
        for col_index, value in enumerate(row):
            if math.isnan(value):
                continue
            label = f"{value:+.2f}"
            text_color = "white" if abs(value) >= vmax * 0.43 else "#202020"
            ax.text(
                col_index,
                row_index,
                label,
                ha="center",
                va="center",
                fontsize=6.2 * FONT_SCALE,
                color=text_color,
                fontfamily="Times New Roman",
            )

    divider = make_axes_locatable(ax)
    cax = divider.append_axes("right", size="3.8%", pad=0.12)
    cbar = fig.colorbar(image, cax=cax)
    cbar.set_label("Completion Rate Delta (skill - base)", fontsize=10 * FONT_SCALE)
    cbar.ax.tick_params(labelsize=9 * FONT_SCALE)
    cbar.set_ticks([-vmax, -vmax / 2, 0.0, vmax / 2, vmax])

    ax.set_xlabel("Agent&Model", labelpad=12, fontsize=AXIS_TITLE_FONT_SIZE)
    ax.set_ylabel("Task Family Domains", labelpad=8, fontsize=AXIS_TITLE_FONT_SIZE)

    for spine in ax.spines.values():
        spine.set_visible(False)

    fig.tight_layout(pad=0.6)
    written_paths: list[Path] = []
    for output_path in output_paths:
        fig.savefig(output_path, bbox_inches="tight", facecolor="white")
        written_paths.append(output_path)
    plt.close(fig)
    return written_paths


def write_long_csv(cells: list[CellValue], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([
            "family_label",
            "agent_model_label",
            "base_slug",
            "skill_slug",
            "n_tasks",
            "delta_completion_rate",
        ])
        for cell in cells:
            writer.writerow([
                cell.family_label,
                cell.agent_model_label,
                cell.base_slug,
                cell.skill_slug,
                cell.n_tasks,
                "" if math.isnan(cell.delta) else f"{cell.delta:+.4f}",
            ])


def write_matrix_csv(
    family_labels: list[str],
    agent_model_labels: list[str],
    matrix: list[list[float]],
    output_path: Path,
) -> None:
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["family_label", *agent_model_labels])
        for family_label, row in zip(family_labels, matrix):
            writer.writerow([
                family_label,
                *["" if math.isnan(value) else f"{value:+.4f}" for value in row],
            ])


def main() -> int:
    args = parse_args()
    jobs_dir = args.jobs_dir.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    workflow_xlsx = args.workflow_xlsx.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    family_map = load_task_family_map(workflow_xlsx)
    pairs = build_pairs(jobs_dir)
    family_labels, agent_model_labels, matrix, cells = build_matrix(pairs, family_map)

    svg_path = output_dir / f"{args.basename}.svg"
    png_path = output_dir / f"{args.basename}.png"
    pdf_path = output_dir / f"{args.basename}.pdf"
    csv_path = output_dir / f"{args.basename}.csv"
    matrix_csv_path = output_dir / f"{args.basename}_matrix.csv"

    render_heatmap(family_labels, agent_model_labels, matrix, [svg_path, png_path, pdf_path], args.vmax)
    write_long_csv(cells, csv_path)
    write_matrix_csv(family_labels, agent_model_labels, matrix, matrix_csv_path)

    print(f"[OK] Wrote SVG: {svg_path}")
    print(f"[OK] Wrote PNG : {png_path}")
    print(f"[OK] Wrote PDF : {pdf_path}")
    print(f"[OK] Wrote CSV: {csv_path}")
    print(f"[OK] Wrote matrix CSV: {matrix_csv_path}")
    print(f"[OK] Task family labels loaded: {len(family_map)}")
    print(f"[OK] Matched model pairs: {len(pairs)}")
    print(f"[OK] Task families: {len(family_labels)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
