#!/usr/bin/env python3
"""Plot an all-task heatmap for skill completion-rate gains.

This variant keeps every individual task/workflow as a separate row while using
``Agent + Model`` on the x-axis. Cells show
``completion_rate(skill) - completion_rate(base)``.
"""

from __future__ import annotations

import argparse
import csv
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from plot_skill_gain_heatmap import (
    DEFAULT_JOBS_DIR,
    DEFAULT_OUTPUT_DIR,
    DEFAULT_VMAX,
    DEFAULT_WORKFLOW_XLSX,
    build_pairs,
    normalize_header,
    pick_english_label,
    wrap_agent_model_label,
)

DEFAULT_BASENAME = "jobs_skill_gain_heatmap_all_tasks_completion_delta"
FONT_SCALE = 1.8
AXIS_TITLE_FONT_SIZE = 30.0
CELL_SIZE_IN = 0.8
FIG_MARGIN_X = 4.2
FIG_MARGIN_Y = 3.6
CELL_BORDER_WIDTH = 1.1

TASK_LABEL_FALLBACKS = {
    "econ-detrending-correlation": "Industry Correlation Analysis",
    "exceltable-in-ppt": "Embedded Data Repair",
    "invoice-fraud-detection": "Document Fraud Detection",
    "jpg-ocr-stat": "OCR Data Extraction",
    "lab-unit-harmonization": "Medical Data Standardization",
    "pptx-reference-formatting": "PPT Formatting Optimization",
    "sales-pivot-analysis": "Sales Pivot Analysis",
    "sec-financial-report": "SEC 13F Financial Analysis",
}


@dataclass(frozen=True)
class TaskCellValue:
    task_name: str
    workflow_label: str
    family_label: str
    agent_model_label: str
    base_slug: str
    skill_slug: str
    delta: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Plot an all-task skill-gain heatmap from Harbor job summaries.")
    parser.add_argument(
        "--jobs-dir",
        type=Path,
        default=DEFAULT_JOBS_DIR,
        help=f"Directory containing top-level jobs folders (default: {DEFAULT_JOBS_DIR})",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for figure outputs (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument(
        "--workflow-xlsx",
        type=Path,
        default=DEFAULT_WORKFLOW_XLSX,
        help=f"Excel lookup table for workflow names (default: {DEFAULT_WORKFLOW_XLSX})",
    )
    parser.add_argument(
        "--basename",
        default=DEFAULT_BASENAME,
        help="Base filename for generated heatmap files.",
    )
    parser.add_argument(
        "--vmax",
        type=float,
        default=DEFAULT_VMAX,
        help="Symmetric color limit for completion deltas (default: 0.4).",
    )
    return parser.parse_args()


def _fallback_workflow_label(task_name: str) -> str:
    if task_name in TASK_LABEL_FALLBACKS:
        return TASK_LABEL_FALLBACKS[task_name]
    if task_name.startswith("harbor_gdpval_"):
        suffix = task_name.removeprefix("harbor_gdpval_")
        return f"GDP Value #{suffix}"
    if task_name.startswith("merge_"):
        suffix = task_name.removeprefix("merge_")
        return f"Merge {suffix.replace('_', '/')}"
    return task_name.replace("_", " ").replace("-", " ")


def _compact_family_label(value: str | None) -> str:
    if not value:
        return "Other"
    compact = value.replace(" and ", " & ")
    return compact


def load_task_metadata(xlsx_path: Path) -> tuple[list[str], dict[str, str], dict[str, str]]:
    task_order: list[str] = []
    workflow_labels: dict[str, str] = {}
    family_labels: dict[str, str] = {}
    if not xlsx_path.is_file():
        return task_order, workflow_labels, family_labels
    try:
        import openpyxl
    except Exception:
        return task_order, workflow_labels, family_labels

    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return task_order, workflow_labels, family_labels

    headers = [normalize_header(value) for value in rows[0]]
    workflow_name_index = next((i for i, h in enumerate(headers) if "workflow name" in h), None)
    core_category_index = next((i for i, h in enumerate(headers) if "core category" in h), None)
    concise_definition_index = next((i for i, h in enumerate(headers) if "concise definition" in h), None)
    if workflow_name_index is None:
        return task_order, workflow_labels, family_labels

    current_family: str | None = None
    for row in rows[1:]:
        raw_task_name = row[workflow_name_index]
        if raw_task_name is None:
            continue
        task_name = str(raw_task_name).strip()
        if not task_name:
            continue
        if task_name not in task_order:
            task_order.append(task_name)
        if core_category_index is not None and row[core_category_index] is not None:
            current_family = pick_english_label(row[core_category_index])
        family_labels[task_name] = _compact_family_label(current_family)
        if concise_definition_index is not None:
            english_label = pick_english_label(row[concise_definition_index])
            if english_label:
                workflow_labels[task_name] = english_label
    return task_order, workflow_labels, family_labels


def wrap_workflow_label(label: str) -> str:
    label = label.replace(" & ", " &\n")
    label = label.replace(" / ", "/\n")
    if len(label) <= 28:
        return label
    words = label.split()
    if len(words) <= 2:
        return label
    midpoint = math.ceil(len(words) / 2)
    return " ".join(words[:midpoint]) + "\n" + " ".join(words[midpoint:])


def build_all_task_matrix(
    jobs_dir: Path,
    workflow_task_order: list[str],
    workflow_labels: dict[str, str],
    family_labels: dict[str, str],
):
    pairs = build_pairs(jobs_dir)
    agent_model_labels = [pair.agent_model_label for pair in pairs]

    available_tasks = set()
    for pair in pairs:
        available_tasks.update(task for task in pair.task_order if task in pair.base_rates and task in pair.skill_rates)

    ordered_tasks = [task for task in workflow_task_order if task in available_tasks]
    remaining_tasks = sorted(available_tasks - set(ordered_tasks))
    ordered_tasks.extend(remaining_tasks)

    matrix: list[list[float]] = []
    cells: list[TaskCellValue] = []
    display_labels: list[str] = []
    for task_name in ordered_tasks:
        workflow_label = workflow_labels.get(task_name, _fallback_workflow_label(task_name))
        family_label = family_labels.get(task_name, "Other")
        display_labels.append(workflow_label)
        row: list[float] = []
        for pair in pairs:
            base_rate = pair.base_rates.get(task_name, math.nan)
            skill_rate = pair.skill_rates.get(task_name, math.nan)
            delta = skill_rate - base_rate if not math.isnan(base_rate) and not math.isnan(skill_rate) else math.nan
            row.append(delta)
            cells.append(
                TaskCellValue(
                    task_name=task_name,
                    workflow_label=workflow_label,
                    family_label=family_label,
                    agent_model_label=pair.agent_model_label,
                    base_slug=pair.base_slug,
                    skill_slug=pair.skill_slug,
                    delta=delta,
                )
            )
        matrix.append(row)
    return ordered_tasks, display_labels, agent_model_labels, matrix, cells


def render_heatmap(
    workflow_labels: list[str],
    agent_model_labels: list[str],
    matrix: list[list[float]],
    output_paths: list[Path],
    vmax: float,
) -> list[Path]:
    try:
        import matplotlib.pyplot as plt
        import numpy as np
        from mpl_toolkits.axes_grid1 import make_axes_locatable
    except Exception as exc:
        msg = "matplotlib and numpy are required to render the heatmap"
        raise RuntimeError(msg) from exc

    data = np.array(matrix, dtype=float)
    masked = np.ma.masked_invalid(data)

    plt.rcParams.update(
        {
            "font.family": "Times New Roman",
            "svg.fonttype": "none",
            "font.size": 9 * FONT_SCALE,
            "axes.labelsize": 11 * FONT_SCALE,
            "xtick.labelsize": 8.8,
            "ytick.labelsize": 8.5 * FONT_SCALE,
        }
    )

    fig_width = max(12.0, CELL_SIZE_IN * len(agent_model_labels) + FIG_MARGIN_X)
    fig_height = max(12.0, CELL_SIZE_IN * len(workflow_labels) + FIG_MARGIN_Y)
    fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=240)

    cmap = plt.get_cmap("RdBu_r").copy()
    cmap.set_bad("#F3F4F6")
    image = ax.imshow(
        masked,
        cmap=cmap,
        vmin=-vmax,
        vmax=vmax,
        aspect="equal",
        interpolation="nearest",
    )

    ax.set_xticks(range(len(agent_model_labels)))
    ax.set_xticklabels(
        [wrap_agent_model_label(label) for label in agent_model_labels],
        rotation=18,
        ha="right",
        rotation_mode="anchor",
    )
    ax.set_yticks(range(len(workflow_labels)))
    ax.set_yticklabels([wrap_workflow_label(label) for label in workflow_labels])
    ax.tick_params(axis="x", length=0, pad=10)
    ax.tick_params(axis="y", length=0, pad=6)

    ax.set_xticks([index - 0.5 for index in range(len(agent_model_labels) + 1)], minor=True)
    ax.set_yticks([index - 0.5 for index in range(len(workflow_labels) + 1)], minor=True)
    ax.grid(which="minor", color="white", linewidth=CELL_BORDER_WIDTH)
    ax.tick_params(which="minor", bottom=False, left=False)

    for row_index, row in enumerate(matrix):
        for col_index, value in enumerate(row):
            if math.isnan(value):
                continue
            label = f"{value:+.2f}"
            text_color = "white" if abs(value) >= vmax * 0.43 else "#202020"
            ax.text(
                col_index,
                row_index,
                label,
                ha="center",
                va="center",
                fontsize=5.6 * FONT_SCALE,
                color=text_color,
                fontfamily="Times New Roman",
            )

    divider = make_axes_locatable(ax)
    cax = divider.append_axes("right", size="3.8%", pad=0.12)
    cbar = fig.colorbar(image, cax=cax)
    cbar.set_label("Completion Rate Delta (skill - base)", fontsize=10 * FONT_SCALE)
    cbar.ax.tick_params(labelsize=9 * FONT_SCALE)
    cbar.set_ticks([-vmax, -vmax / 2, 0.0, vmax / 2, vmax])

    ax.set_xlabel("Agent&Model", labelpad=12, fontsize=AXIS_TITLE_FONT_SIZE)
    ax.set_ylabel("All Tasks Family", labelpad=8, fontsize=AXIS_TITLE_FONT_SIZE)

    for spine in ax.spines.values():
        spine.set_visible(False)

    fig.tight_layout(pad=0.7)
    written_paths: list[Path] = []
    for output_path in output_paths:
        fig.savefig(output_path, bbox_inches="tight", facecolor="white")
        written_paths.append(output_path)
    plt.close(fig)
    return written_paths


def write_long_csv(cells: list[TaskCellValue], output_path: Path) -> None:
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "task_name",
                "workflow_label",
                "family_label",
                "agent_model_label",
                "base_slug",
                "skill_slug",
                "delta_completion_rate",
            ]
        )
        for cell in cells:
            writer.writerow(
                [
                    cell.task_name,
                    cell.workflow_label,
                    cell.family_label,
                    cell.agent_model_label,
                    cell.base_slug,
                    cell.skill_slug,
                    "" if math.isnan(cell.delta) else f"{cell.delta:+.4f}",
                ]
            )


def write_matrix_csv(
    task_names: list[str],
    workflow_labels: list[str],
    family_labels: dict[str, str],
    agent_model_labels: list[str],
    matrix: list[list[float]],
    output_path: Path,
) -> None:
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["task_name", "workflow_label", "family_label", *agent_model_labels])
        for task_name, workflow_label, row in zip(task_names, workflow_labels, matrix):
            writer.writerow(
                [
                    task_name,
                    workflow_label,
                    family_labels.get(task_name, "Other"),
                    *["" if math.isnan(value) else f"{value:+.4f}" for value in row],
                ]
            )


def main() -> int:
    args = parse_args()
    output_dir = args.output_dir.expanduser().resolve()
    workflow_xlsx = args.workflow_xlsx.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    task_order, workflow_label_map, family_label_map = load_task_metadata(workflow_xlsx)
    jobs_dir = args.jobs_dir.expanduser().resolve()
    task_names, workflow_labels, agent_model_labels, matrix, cells = build_all_task_matrix(
        jobs_dir,
        task_order,
        workflow_label_map,
        family_label_map,
    )

    svg_path = output_dir / f"{args.basename}.svg"
    png_path = output_dir / f"{args.basename}.png"
    pdf_path = output_dir / f"{args.basename}.pdf"
    csv_path = output_dir / f"{args.basename}.csv"
    matrix_csv_path = output_dir / f"{args.basename}_matrix.csv"

    render_heatmap(workflow_labels, agent_model_labels, matrix, [svg_path, png_path, pdf_path], args.vmax)
    write_long_csv(cells, csv_path)
    write_matrix_csv(task_names, workflow_labels, family_label_map, agent_model_labels, matrix, matrix_csv_path)

    print(f"[OK] Wrote SVG: {svg_path}")
    print(f"[OK] Wrote PNG : {png_path}")
    print(f"[OK] Wrote PDF : {pdf_path}")
    print(f"[OK] Wrote CSV: {csv_path}")
    print(f"[OK] Wrote matrix CSV: {matrix_csv_path}")
    print(f"[OK] Tasks in heatmap: {len(task_names)}")
    print(f"[OK] Matched model pairs: {len(agent_model_labels)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
